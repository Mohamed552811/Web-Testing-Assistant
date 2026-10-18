import os
import tempfile
import unittest
import zipfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
if os.name == "nt":  # the offscreen platform has no fonts of its own; text would render as boxes
    os.environ.setdefault("QT_QPA_FONTDIR", os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"))

from PySide6.QtGui import QColor, QImage  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from wta import settings as app_settings  # noqa: E402
from wta.db import Database, store  # noqa: E402
from wta.db import scenarios as sc  # noqa: E402
from wta.db.models import Bug, Screenshot, Step, StepStatus  # noqa: E402
from wta.export import build_report, export_report  # noqa: E402
from wta.export.xlsx import png_size  # noqa: E402
from wta.prompt import RunContext, SmokeContext, build_prompt  # noqa: E402
from wta.results import StepResult, screenshot_rule  # noqa: E402
from wta.settings_dialog import SettingsDialog  # noqa: E402

app = QApplication.instance() or QApplication([])
PASSWORD = "Secret-Pw-123"


def make_png(path: Path, width=200, height=120, color="red") -> Path:
    image = QImage(width, height, QImage.Format.Format_RGB32)
    image.fill(QColor(color))
    image.save(str(path), "PNG")
    return path


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.base = Path(cls.tmp.name)
        cls.db = Database(cls.base / "wta.sqlite3")
        session = store.create_session(cls.db, store.SessionDetails(
            name="Shop | checkout", target_url="https://shop.example.com", test_account_email="qa@shop.example.com",
            application="Shop web app", test_scope="Checkout\n  \nCart",
            test_account_password=PASSWORD))
        cls.session_id = session.id
        bug_png = make_png(cls.base / "bug.png")
        blocked_png = make_png(cls.base / "blocked.png", color="yellow")

        scenario = sc.add_scenario(cls.db, session.id, 1, sc.ScenarioInput(
            "Checkout", "Open shop\nPay for one item", "The order is placed", "Card: test card"))
        smoke = sc.add_smoke_page(cls.db, session.id, 1, sc.SmokePageInput("Cart page", "https://shop.example.com/cart"))
        with cls.db.Session.begin() as orm:
            steps = [
                Step(scenario_id=scenario.id, order=1, action="Open shop", status=StepStatus.PASS,
                     actual_result="ok | fine", reasoning="Loaded"),
                Step(scenario_id=scenario.id, order=2, action="Pay", status=StepStatus.FAIL, actual_result="500 error",
                     bugs=[Bug(title="Payment fails", severity="2 - High", priority=1,
                               repro_steps="1. Sign in as QA@shop.example.com\n2. Add item\n3. Pay",
                               expected="Order placed", actual="500 error",
                               report_text=f"Full report (typed {PASSWORD} by mistake)")],
                     screenshots=[Screenshot(file_path=str(bug_png))]),
                Step(scenario_id=scenario.id, order=3, action="Coupon", status=StepStatus.FAIL,
                     bugs=[Bug(title="False alarm", severity="4 - Low", priority=4)]),
                Step(scenario_id=scenario.id, order=4, action="Track order", status=StepStatus.BLOCKED,
                     screenshots=[Screenshot(file_path=str(blocked_png))]),
                Step(scenario_id=smoke.id, order=1, action="Empty cart message", status=StepStatus.FAIL,
                     bugs=[Bug(title="No empty-cart message", severity="3 - Medium", priority=2)]),
            ]
            orm.add_all(steps)
        ids = [s.id for s in steps]
        sc.approve_step(cls.db, ids[0], "fine")
        sc.approve_step(cls.db, ids[1], "real bug")
        sc.override_step(cls.db, ids[2], StepStatus.PASS, "coupon works, Claude misread")
        cls.report = build_report(cls.db, session.id)

    @classmethod
    def tearDownClass(cls):
        cls.db.close()
        cls.tmp.cleanup()

    # model

    def test_report_model(self):
        r = self.report
        self.assertEqual([(row.kind, row.order) for row in r.rows],
                         [("Scenario", 1), ("Scenario", 2), ("Scenario", 3), ("Scenario", 4), ("Smoke test", 1)])
        self.assertEqual(r.rows[2].claude_status, "Fail")
        self.assertEqual(r.rows[2].final_status, "Pass")  # the override wins, Claude's finding is kept
        self.assertEqual(r.rows[4].claude_status, "Bug")  # smoke label
        self.assertEqual(r.rows[3].final_status, "Pending review")
        self.assertEqual(r.counts["Scenario"]["Pass"], 2)
        self.assertEqual(r.counts["Scenario"]["Fail / Bug"], 1)
        self.assertEqual(r.counts["Scenario"]["Pending review"], 1)
        # Overridden-to-Pass bug left out; unreviewed smoke bug kept.
        self.assertEqual([(b.title, b.review_state) for b in r.bugs],
                         [("Payment fails", "Confirmed"), ("No empty-cart message", "Not reviewed yet")])
        self.assertIn("Expected: Order placed", r.bugs[0].tfs_repro_steps)
        # The password never leaves the machine in a report; the account email may (the user's reports name it).
        self.assertIn("1. Sign in as QA@shop.example.com", r.bugs[0].repro_steps)
        self.assertEqual(r.bugs[0].report_text, "Full report (typed the test password by mistake)")
        self.assertEqual(r.test_account_email, "qa@shop.example.com")
        self.assertEqual(len(r.bugs[0].screenshots), 1)

    def no_secrets(self, data: bytes, what: str):
        self.assertNotIn(PASSWORD.encode(), data, what)

    # formats

    def test_xlsx(self):
        from openpyxl import load_workbook
        path = export_report(self.db, self.session_id, self.base / "out" / "report.xlsx")
        wb = load_workbook(path)
        self.assertEqual(wb.sheetnames, ["Summary", "Results", "Bugs"])
        bugs = wb["Bugs"]
        headers = [c.value for c in bugs[1]]
        for name in ("Work Item Type", "Title", "Repro Steps", "Severity", "Priority", "Tags"):
            self.assertIn(name, headers)
        self.assertEqual(bugs.max_row, 3)
        self.assertEqual(bugs["A2"].value, "Bug")
        self.assertEqual(bugs["D2"].value, "2 - High")
        self.assertEqual(wb["Results"].max_row, 6)
        with zipfile.ZipFile(path) as z:
            media = [n for n in z.namelist() if n.startswith("xl/media/")]
            self.assertEqual(len(media), 4)  # logo + bug and blocked in Results + bug again in Bugs
            self.assertEqual(z.read(media[0])[:8], b"\x89PNG\r\n\x1a\n")
            for name in z.namelist():
                self.no_secrets(z.read(name), name)

    def test_docx(self):
        from docx import Document
        path = export_report(self.db, self.session_id, self.base / "out" / "report", "docx")
        self.assertEqual(path.suffix, ".docx")
        doc = Document(str(path))
        self.assertEqual(len(doc.inline_shapes), 3)  # logo + bug screenshot + blocked screenshot
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("1. Payment fails", text)
        self.assertIn("Other screenshots", text)
        with zipfile.ZipFile(path) as z:
            for name in z.namelist():
                self.no_secrets(z.read(name), name)

    def test_pdf(self):
        from PySide6.QtPdf import QPdfDocument
        path = export_report(self.db, self.session_id, self.base / "out" / "report.pdf")
        data = path.read_bytes()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 5000)
        document = QPdfDocument()
        self.assertEqual(document.load(str(path)), QPdfDocument.Error.None_)
        self.assertGreaterEqual(document.pageCount(), 1)
        text = "".join(document.getAllText(i).text() for i in range(document.pageCount()))
        self.assertIn("Payment fails", text)
        self.assertNotIn(PASSWORD, text)
        document.close()

    def test_markdown_summary_and_bug_files(self):
        summary = export_report(self.db, self.session_id, self.base / "out" / "report.md", tester="Staff QA")
        folder = self.base / "out" / "report"
        self.assertEqual(summary, folder / "SUMMARY.md")
        self.assertEqual(sorted(p.name for p in folder.iterdir()), [
            "BUG-001-payment-fails.md", "BUG-001-screenshot.png", "BUG-002-no-empty-cart-message.md", "SUMMARY.md",
            "TC001-step-4.png", "logo.png",
        ])
        self.assertTrue(summary.read_text("utf-8").startswith('<img src="logo.png"'))
        text = summary.read_text("utf-8")
        for needle in (
            "# Shop | checkout Test Report", "## Test Information", "- **Tester:** Staff QA",
            "- **Application:** Shop web app", "- **URL:** https://shop.example.com",
            "- **User:** qa@shop.example.com", "## Test Scope\n- Checkout\n- Cart",
            "### Test Case 001: Checkout", "| Test Step | Expected Result | Actual Result | Status |",
            "| Open shop |  | ok \\| fine | PASS |", "| Pay |  | 500 error | FAIL |",
            "| Coupon |  |  | PASS |",  # the override wins
            "| Track order |  |  | BLOCKED (not reviewed) |",
            "### Smoke Test 002: Cart page", "| Empty cart message |  |  | FAIL (not reviewed) |",
            "## Overall Test Result\n**FAIL (2 steps not reviewed yet)**",
            "## Key Findings\n\n## Bugs Found",
            "| [BUG-001](BUG-001-payment-fails.md) | Payment fails | High | Open |",
            "| [BUG-002](BUG-002-no-empty-cart-message.md) | No empty-cart message | Medium | Open (not reviewed) |",
            "## Recommendations\n\n## Attachments",
            "- BUG-001-screenshot.png - Checkout, step 2: 500 error",
            "- TC001-step-4.png - Checkout, step 4: Track order",
        ):
            self.assertIn(needle, text)
        self.assertNotIn("False alarm", text)  # overridden to Pass: no bug

        bug = (folder / "BUG-001-payment-fails.md").read_text("utf-8")
        for needle in (
            "# BUG-001: Payment fails", "## Severity: High",
            "## Summary\nFull report (typed the test password by mistake)",
            "## Environment\n- **URL:** https://shop.example.com", "- **User:** qa@shop.example.com",
            "## Steps to Reproduce\n1. Sign in as QA@shop.example.com\n2. Add item\n3. Pay",
            "## Expected Behavior\n- Order placed", "## Actual Behavior\n- 500 error",
            "## Screenshots\n- See: BUG-001-screenshot.png (Checkout, step 2: 500 error)",
            "## Impact\n\n## Notes\n- Review: real bug",
        ):
            self.assertIn(needle, bug)
        for f in folder.glob("*.md"):
            self.no_secrets(f.read_bytes(), f.name)

    def test_markdown_slug_and_reexport(self):
        from wta.export.markdown import severity_word, slug
        self.assertEqual(slug('Price List "strawberry export pl 2" Not Found in New Sales Order'),
                         "price-list-strawberry-export-pl-2-not-found-in-new")
        self.assertEqual(slug("!!!"), "bug")
        self.assertEqual(severity_word("2 - High"), "High")
        # Exporting again to the same name replaces the folder instead of mixing old files in.
        target = self.base / "again" / "r.md"
        export_report(self.db, self.session_id, target)
        (target.with_suffix("") / "stale.txt").write_text("old")
        export_report(self.db, self.session_id, target)
        self.assertFalse((target.with_suffix("") / "stale.txt").exists())

    def test_unknown_format(self):
        with self.assertRaises(ValueError):
            export_report(self.db, self.session_id, self.base / "x.txt")

    def test_png_size(self):
        self.assertEqual(png_size(make_png(self.base / "size.png", 321, 123)), (321, 123))
        (self.base / "not.png").write_bytes(b"GIF89a....................")
        self.assertIsNone(png_size(self.base / "not.png"))


class ScreenshotTimingTests(unittest.TestCase):
    def step(self, order, status):
        return StepResult(order, "", "", "a", "", "", status, "", "", [], [], [], None, None)

    def test_rule(self):
        default = app_settings.SCREENSHOTS_BUG_BLOCKED
        every = app_settings.SCREENSHOTS_EVERY_STEP
        once = app_settings.SCREENSHOTS_ONCE_PER_PAGE
        self.assertEqual(screenshot_rule(default, False, self.step(1, StepStatus.PASS)), (False, False))
        self.assertEqual(screenshot_rule(default, False, self.step(2, StepStatus.FAIL)), (True, True))
        self.assertEqual(screenshot_rule(every, False, self.step(2, StepStatus.PASS)), (True, True))
        self.assertEqual(screenshot_rule(once, True, self.step(1, StepStatus.PASS)), (True, True))
        self.assertEqual(screenshot_rule(once, True, self.step(2, StepStatus.PASS)), (False, False))
        self.assertEqual(screenshot_rule(once, False, self.step(1, StepStatus.PASS)), (False, False))

    def test_prompt_follows_timing(self):
        ctx = RunContext("S", "https://a.example.com", ["a.example.com"], "", "", "N", "G", "")
        self.assertIn('Screenshots ONLY for "fail" and "blocked"', build_prompt(ctx))
        ctx.screenshot_timing = app_settings.SCREENSHOTS_EVERY_STEP
        self.assertIn("Screenshots for EVERY step", build_prompt(ctx))
        ctx.screenshot_timing = app_settings.SCREENSHOTS_ONCE_PER_PAGE
        self.assertIn('Screenshots ONLY for "fail"', build_prompt(ctx))  # scenarios: like the default
        ctx.smoke = SmokeContext("P", "https://a.example.com/p")
        self.assertIn("Screenshots for step 1 (one view of the page)", build_prompt(ctx))


class SettingsTests(unittest.TestCase):
    def test_load_save_and_clean(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            s = app_settings.load(path)
            self.assertEqual((s.screenshot_timing, s.export_format, s.confirm_risky_actions),
                             ("bug_blocked", "xlsx", False))
            app_settings.save(path, app_settings.Settings("every_step", "pdf", True))
            self.assertEqual(app_settings.load(path), app_settings.Settings("every_step", "pdf", True))
            path.write_text('{"screenshot_timing": "sometimes", "export_format": "odt"}', "utf-8")
            self.assertEqual(app_settings.load(path), app_settings.Settings())

    def test_tester_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            app_settings.save(path, app_settings.Settings(tester_name="  Staff QA  "))
            self.assertEqual(app_settings.load(path).tester_name, "Staff QA")

    def test_session_application_and_scope(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(Path(tmp) / "w.sqlite3")
            try:
                s = store.create_session(db, store.SessionDetails(
                    name="S", target_url="https://a.example.com", application=" ERP ", test_scope="One\n\n Two \n"))
                details = store.details_of(store.load_session(db, s.id))
                self.assertEqual((details.application, details.test_scope), ("ERP", "One\nTwo"))
            finally:
                db.close()

    def test_dialog_round_trip(self):
        dialog = SettingsDialog(app_settings.Settings("once_per_smoke_page", "docx", True, "Staff QA"))
        self.assertEqual(dialog.settings(), app_settings.Settings("once_per_smoke_page", "docx", True, "Staff QA"))
        dialog.export_format.setCurrentIndex(dialog.export_format.findData("md"))
        self.assertEqual(dialog.settings().export_format, "md")
        dialog.deleteLater()


if __name__ == "__main__":
    unittest.main()


class ScenarioStepsInReportsTests(ExportTests):
    """Phase 11: every format shows the scenario's steps and expected result."""

    def test_model(self):
        (info,) = self.report.scenarios  # the smoke page is not a scenario
        self.assertEqual((info.name, info.steps, info.expected_result, info.test_data),
                         ("Checkout", ["Open shop", "Pay for one item"], "The order is placed", "Card: test card"))
        self.assertEqual(info.numbered_steps(), "1. Open shop\n2. Pay for one item")

    def test_markdown_lists_steps_and_expected_under_the_test_case(self):
        text = export_report(self.db, self.session_id, self.base / "steps" / "r.md").read_text("utf-8")
        self.assertIn("### Test Case 001: Checkout\n\n**Steps:**\n1. Open shop\n2. Pay for one item\n\n"
                      "**Expected result:** The order is placed\n\n**Test data:** Card: test card\n\n| Test Step", text)

    def test_excel_summary_has_scenarios_table(self):
        from openpyxl import load_workbook
        ws = load_workbook(export_report(self.db, self.session_id, self.base / "steps" / "r.xlsx"))["Summary"]
        values = [[c.value for c in row] for row in ws.iter_rows()]
        header = next(i for i, r in enumerate(values) if r[:5] == ["Batch", "Scenario", "Steps", "Expected result", "Test data"])
        self.assertEqual(values[header + 1][:5],
                         [1, "Checkout", "1. Open shop\n2. Pay for one item", "The order is placed", "Card: test card"])

    def test_word_and_pdf_have_scenarios_section(self):
        from docx import Document
        doc = Document(str(export_report(self.db, self.session_id, self.base / "steps" / "r.docx")))
        self.assertIn("Scenarios", [p.text for p in doc.paragraphs])
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        self.assertIn("1. Open shop\n2. Pay for one item", cells)
        self.assertIn("The order is placed", cells)
        from wta.export.html import report_html
        html = report_html(self.report)
        self.assertIn("<h2>Scenarios</h2>", html)
        self.assertIn("1. Open shop<br>2. Pay for one item", html)
