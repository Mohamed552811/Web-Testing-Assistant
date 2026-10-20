"""M6: Send to Smoke / Regression (checklist runs), Add as scenarios and Copy as prompt (spec 8.6, decision 2)."""

import csv
import json
import os
import re
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
if os.name == "nt":  # the offscreen platform has no fonts of its own; text would render as boxes
    os.environ.setdefault("QT_QPA_FONTDIR", os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"))

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtGui import QGuiApplication  # noqa: E402
from PySide6.QtWidgets import QDialog, QFileDialog, QMessageBox  # noqa: E402

from wta import launchpad as modes  # noqa: E402
from wta.app import MainWindow  # noqa: E402
from wta.checklist_tab import ChecklistTab, case_html  # noqa: E402
from wta.db import Database, checklists, library  # noqa: E402
from wta.db import designer_content as content  # noqa: E402
from wta.db.models import ScenarioKind  # noqa: E402
from wta.designer import handoff  # noqa: E402
from wta.designer.handoff_dialog import PROMPT, SCENARIOS, HandOffDialog  # noqa: E402
from wta.designer.settings import DesignerSettings  # noqa: E402
from wta.designer.workspace import DesignerWorkspace  # noqa: E402

try:
    from qt_app import app  # noqa: E402
    from test_designer_generation import case
    from test_designer_history_export import Base
except ImportError:  # run as tests.<module>
    from tests.qt_app import app  # noqa: E402
    from tests.test_designer_generation import case
    from tests.test_designer_history_export import Base

SPEC = Path(__file__).resolve().parent.parent / "docs" / "test-case-designer" / "TEST_CASE_DESIGNER_SPEC.md"


class HandOffBase(Base):
    def setUp(self):
        super().setUp()
        self.save(self.orders.id, case(1), case(2, "corner_case"), case(3, "corner_case"))
        self.cases = content.current_cases(self.db, self.orders.id)
        self.attachments = self.dir / "checklists"

    def handoff_cases(self, cases=None):
        return [checklists.HandOffCase(c.row_id, self.orders.id, c.data) for c in (cases or self.cases)]


class ChecklistStoreTests(HandOffBase):
    def test_run_lifecycle(self):
        self.assertEqual(checklists.default_name("smoke", "Orders v1.xlsx", datetime(2026, 10, 20)),
                         "Smoke – Orders v1 – 2026-10-20")
        run_id = checklists.create_run(self.db, "smoke", "Smoke – Orders", self.handoff_cases(self.cases[:2]),
                                       {"workbook": "Orders v1.xlsx"})
        self.assertEqual(checklists.open_count(self.db, "smoke"), 1)
        self.assertEqual(checklists.open_count(self.db, "regression"), 0)
        run = checklists.load_run(self.db, run_id)
        self.assertEqual((run.status, [i.position for i in run.items]), ("draft", [1, 2]))
        self.assertEqual(run.items[0].snapshot["id"], "TC-007-01")

        # Later designer edits do not change the run, until a refresh while it is still a draft.
        content.update_case(self.db, self.cases[0].row_id, dict(self.cases[0].data, title="Verify that it changed"))
        self.assertEqual(checklists.load_run(self.db, run_id).items[0].snapshot["title"], "Verify that flow number 1 works")
        self.assertEqual(checklists.refresh_from_designer(self.db, run_id), 1)
        self.assertEqual(checklists.load_run(self.db, run_id).items[0].snapshot["title"], "Verify that it changed")

        self.assertEqual(checklists.add_to_run(self.db, run_id, self.handoff_cases()), 1)  # only TC-007-03 is new
        first, second, third = checklists.load_run(self.db, run_id).items
        checklists.record(self.db, first.id, "pass")
        checklists.record(self.db, second.id, "fail", "Export button missing")
        shot = self.dir / "shot.png"
        shot.write_bytes(b"\x89PNG fake")
        relative = checklists.attach(self.db, second.id, shot, self.attachments)
        self.assertTrue((self.attachments / relative).is_file())
        run = checklists.load_run(self.db, run_id)
        self.assertEqual(run.status, "in_progress")
        self.assertEqual(run.counts(), {"pass": 1, "fail": 1, "blocked": 0, "skipped": 0, "not_run": 1})
        self.assertEqual(run.progress_text, "2 of 3 done · Pass 1 · Fail 1 · Blocked 0 · Skipped 0")
        self.assertEqual((run.items[1].notes, run.items[1].attachment), ("Export button missing", relative))
        self.assertIsNotNone(run.items[0].executed_at)
        with self.assertRaises(ValueError):
            checklists.refresh_from_designer(self.db, run_id)  # not a draft any more
        with self.assertRaises(ValueError):
            checklists.record(self.db, third.id, "maybe")

        checklists.set_completed(self.db, run_id, True)
        self.assertEqual(checklists.open_count(self.db, "smoke"), 0)
        with self.assertRaises(ValueError):
            checklists.record(self.db, third.id, "pass")
        with self.assertRaises(ValueError):
            checklists.add_to_run(self.db, run_id, self.handoff_cases())
        checklists.set_completed(self.db, run_id, False)
        self.assertEqual(checklists.load_run(self.db, run_id).status, "in_progress")

        run = checklists.load_run(self.db, run_id)
        csv_path = checklists.export_run(run, self.dir / "run.csv", "csv")
        with csv_path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        self.assertEqual(rows[0], list(checklists.EXPORT_COLUMNS))
        self.assertEqual([r[6] for r in rows[1:]], ["Pass", "Fail", "Not run"])
        self.assertTrue(rows[2][4].startswith("1. Open"))
        book = load_workbook(checklists.export_run(run, self.dir / "run.xlsx", "xlsx"))
        self.assertEqual(book.sheetnames, ["Checklist", "Summary"])
        self.assertEqual(book["Checklist"].max_row, 4)

        checklists.delete_run(self.db, run_id, self.attachments)
        self.assertFalse((self.attachments / f"run-{run_id}").exists())
        with self.assertRaises(LookupError):
            checklists.load_run(self.db, run_id)

    def test_validation(self):
        with self.assertRaises(ValueError):
            checklists.create_run(self.db, "smoke", "Empty", [], {})
        with self.assertRaises(ValueError):
            checklists.create_run(self.db, "nightly", "X", self.handoff_cases(), {})


class PromptAndScenarioTests(HandOffBase):
    def test_copy_as_prompt(self):
        text = handoff.copy_as_prompt("smoke", [c.data for c in self.cases], DesignerSettings(product_context="a CRM"))
        self.assertTrue(text.startswith("You are helping a QA tester run a smoke test (manual checklist) for a CRM."))
        cases = json.loads(text.split("Test cases (JSON):\n", 1)[1])
        self.assertEqual([c["id"] for c in cases], ["TC-007-01", "TC-007-02", "TC-007-03"])

    @unittest.skipUnless(SPEC.is_file(), "spec not present")
    def test_template_is_verbatim_spec_8_6(self):
        spec = SPEC.read_text("utf-8")
        section = spec[spec.index("### 8.6"):spec.index("## 9.")]
        expected = re.findall(r"```text\n(.*?)\n```", section, re.S)[0]
        self.assertEqual(handoff.pb._read("copy_as_prompt.txt").rstrip("\n"), expected)

    def test_scenario_input(self):
        item = handoff.scenario_input(case(1, preconditions=["Logged in", "An order exists"]), "Orders v1.xlsx")
        scenario = item.cleaned().scenario
        self.assertEqual(scenario.name, "TC-007-01 Verify that flow number 1 works")
        self.assertIn("Make sure that: Logged in; An order exists", scenario.steps)
        self.assertIn("Assert that the list is shown.", scenario.steps)
        self.assertEqual(scenario.expected_result, "The flow works.")
        per_step = handoff.scenario_input(case(1, steps=[{"action": "Save.", "expected": "Assert that it is saved."}]),
                                          "x")
        self.assertIn("Save. — Assert that it is saved.", per_step.scenario.steps)


class WorkspaceHandOffTests(HandOffBase):
    def setUp(self):
        super().setUp()
        self.workspace = DesignerWorkspace(self.db, self.dir / "runs", None)
        self.workspace.show()
        self.workspace.set_workbook(self.v1.id, "Orders v1.xlsx (v1)")
        self.workspace.select_backlog(self.orders.id)
        self.created = []
        self.workspace.checklist_created.connect(lambda mode, run_id: self.created.append((mode, run_id)))

    def tearDown(self):
        self.workspace.deleteLater()
        super().tearDown()

    def run_handoff(self, kind: str, configure) -> None:
        def accept(dialog):
            configure(dialog)
            return QDialog.DialogCode.Accepted

        with mock.patch.object(HandOffDialog, "exec", accept), \
                mock.patch.object(QMessageBox, "exec", return_value=0), \
                mock.patch.object(QMessageBox, "information"):
            self.workspace.hand_off(kind)

    def test_selected_cases_become_a_smoke_run(self):
        results = self.workspace.results_panel
        results.tag_filter.setCurrentIndex(results.tag_filter.findData("smoke"))
        results.select_shown()
        self.assertEqual([c.data["id"] for c in results.selected_cases()], ["TC-007-01"])
        self.assertEqual(results.selection_label.text(), "1 selected")
        self.run_handoff("smoke", lambda d: None)
        [run] = checklists.list_runs(self.db, "smoke")
        self.assertTrue(run.name.startswith("Smoke – Orders v1 – "))
        self.assertEqual([i.snapshot["id"] for i in run.items], ["TC-007-01"])
        self.assertEqual(run.source["workbook"], "Orders v1.xlsx")

        # Tagged scope, added to the open run.
        def tagged_into_existing(dialog):
            dialog.tagged_radio.setChecked(True)
            dialog.tag_combo.setCurrentIndex(dialog.tag_combo.findData("regression"))
            dialog.existing_radio.setChecked(True)

        self.run_handoff("smoke", tagged_into_existing)
        self.assertEqual(len(checklists.load_run(self.db, run.id).items), 3)

    def test_regression_run_prompt_and_scenarios(self):
        results = self.workspace.results_panel
        results.clear_selection()

        def tagged(dialog):
            dialog.tagged_radio.setChecked(True)
            dialog.tag_combo.setCurrentIndex(dialog.tag_combo.findData("regression"))

        self.run_handoff("regression", tagged)
        [run] = checklists.list_runs(self.db, "regression")
        self.assertEqual(len(run.items), 3)

        def tagged_regression_prompt(dialog):
            tagged(dialog)
            dialog.mode_combo.setCurrentIndex(dialog.mode_combo.findData("regression"))

        self.run_handoff(PROMPT, tagged_regression_prompt)
        clip = QGuiApplication.clipboard().text()
        self.assertIn("run a regression test (manual checklist)", clip)
        self.assertIn("TC-007-03", clip)

        self.run_handoff(SCENARIOS, tagged)
        items = library.list_items(self.db, ScenarioKind.SCENARIO)
        self.assertEqual(len(items), 3)
        self.assertTrue(all(i.source.startswith("Test Case Designer: Orders v1.xlsx") for i in items))

    def test_dialog(self):
        dialog = HandOffDialog("smoke", 0, ["regression", "smoke"], "Smoke – X", [(5, "Old run")],
                               {"smoke": 1, "regression": 3})
        self.assertTrue(dialog.tagged_radio.isChecked())  # nothing selected
        self.assertEqual(dialog.tag(), "smoke")
        dialog.name_edit.clear()
        dialog._accept()
        self.assertIn("name", dialog.error.text())
        dialog.existing_radio.setChecked(True)
        self.assertEqual(dialog.existing_run(), 5)
        prompt = HandOffDialog(PROMPT, 2, [], parent=None)
        prompt.mode_combo.setCurrentIndex(1)
        self.assertEqual((prompt.scope(), prompt.mode()), ("selected", "regression"))
        dialog.deleteLater()
        prompt.deleteLater()


class ChecklistTabTests(HandOffBase):
    def test_tester_flow(self):
        run_id = checklists.create_run(self.db, "regression", "Regression – Orders", self.handoff_cases(), {})
        tab = ChecklistTab(self.db, self.attachments)
        tab.export_dir = self.dir
        tab.show()
        try:
            tab.refresh()
            self.assertEqual(tab.run_combo.currentData(), run_id)
            self.assertEqual(tab.items.count(), 3)
            self.assertIn("TC-007-01", tab.case_view.toPlainText())
            self.assertTrue(tab.refresh_button.isEnabled())
            tab.result_buttons["pass"].click()
            self.assertEqual(tab.items.currentRow(), 1)  # moved on
            tab.notes.setPlainText("Broken export")
            tab.result_buttons["fail"].click()
            run = checklists.load_run(self.db, run_id)
            self.assertEqual([i.result for i in run.items], ["pass", "fail", None])
            self.assertEqual(run.items[1].notes, "Broken export")
            self.assertFalse(tab.refresh_button.isEnabled())  # no longer a draft
            self.assertIn("2 of 3 done", tab.progress.text())
            tab.items.setCurrentRow(1)
            shot = self.dir / "shot.png"
            shot.write_bytes(b"png")
            with mock.patch.object(QFileDialog, "getOpenFileName", return_value=(str(shot), "")):
                tab.attach_screenshot()
            self.assertTrue(tab.open_attachment_button.isEnabled())
            with mock.patch.object(QFileDialog, "getSaveFileName", return_value=(str(self.dir / "r.csv"), "CSV (*.csv)")):
                tab.export()
            self.assertTrue((self.dir / "r.csv").is_file())
            with mock.patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                tab.complete_button.click()
            self.assertEqual(checklists.load_run(self.db, run_id).status, "completed")
            self.assertEqual(tab.complete_button.text(), "Reopen run")
            self.assertFalse(tab.result_buttons["pass"].isEnabled())
            tab.set_mode("smoke")
            self.assertTrue(tab.empty.isVisibleTo(tab))
            self.assertIn("Create Smoke run", tab.empty.text())
        finally:
            tab.deleteLater()

    def test_case_html(self):
        html = case_html(case(1, preconditions=["Logged in"], steps=[{"action": "Save.", "expected": "Assert it."}]))
        self.assertIn("<li>Logged in</li>", html)
        self.assertIn("<i>Expected:</i> Assert it.", html)


class MainWindowTests(HandOffBase):
    def test_open_run_and_badges(self):
        run_id = checklists.create_run(self.db, "smoke", "Smoke – Orders", self.handoff_cases(), {})
        window = MainWindow(self.db, app_dir=self.dir, check_sign_in=False)
        try:
            self.assertEqual(window.launchpad.cards[modes.SMOKE].badge.text(), "1 open checklist run")
            self.assertEqual(window.launchpad.cards[modes.REGRESSION].badge.text(), "")
            window.open_checklist_run("smoke", run_id)
            self.assertEqual(window.mode, modes.SMOKE)
            self.assertIs(window.tabs.currentWidget(), window.checklist_tab)
            self.assertEqual(window.checklist_tab.run_combo.currentData(), run_id)
        finally:
            window.deleteLater()


if __name__ == "__main__":
    unittest.main()
