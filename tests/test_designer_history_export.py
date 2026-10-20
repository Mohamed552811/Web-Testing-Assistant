"""M5: regenerate with locked cases and guidance, history and restore, links across workbook versions, tag filter,
and the JSON / CSV / Excel exports (spec 8.1 (1), (6), (8), (9)). Generation is faked."""

import csv
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
if os.name == "nt":  # the offscreen platform has no fonts of its own; text would render as boxes
    os.environ.setdefault("QT_QPA_FONTDIR", os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"))

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtWidgets import QMessageBox  # noqa: E402
from sqlalchemy import select  # noqa: E402

from wta.db import Database  # noqa: E402
from wta.db import designer_content as content  # noqa: E402
from wta.db import generations  # noqa: E402
from wta.db import workbooks as wb  # noqa: E402
from wta.db.designer_models import ExportLog  # noqa: E402
from wta.designer import export as exporter  # noqa: E402
from wta.designer import workspace as workspace_module  # noqa: E402
from wta.designer.export_dialog import ExportDialog  # noqa: E402
from wta.designer.generation import CallResult, Planner  # noqa: E402
from wta.designer.history_dialog import HistoryDialog  # noqa: E402
from wta.designer.settings import PER_STEP_EXPECTED, DesignerSettings, OutputPreset  # noqa: E402
from wta.designer.workspace import DesignerWorkspace  # noqa: E402

try:
    from designer_xlsx import Pic, SheetSpec, build, png
    from qt_app import app  # noqa: E402
    from test_designer_generation import case, response
    from test_designer_workspace import FakeRun
except ImportError:  # run as tests.<module>
    from tests.designer_xlsx import Pic, SheetSpec, build, png
    from tests.qt_app import app  # noqa: E402
    from tests.test_designer_generation import case, response
    from tests.test_designer_workspace import FakeRun


def sheets(text: str = "Orders can be exported.") -> list[SheetSpec]:
    return [SheetSpec("Main", {"A1": 1, "B1": "Login", "B2": "The user signs in.", "A4": 7, "B4": "Orders", "B5": text},
                      [Pic("B6", png(seed=2))])]


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.db = Database(self.dir / "wta.sqlite3")
        self.v1 = self.upload("Orders v1.xlsx")
        self.orders = self.v1.sheets[0].backlogs[1]

    def tearDown(self):
        app.processEvents()
        self.db.close()
        self.tmp.cleanup()

    def upload(self, name: str, text: str = "Orders can be exported.") -> wb.WorkbookView:
        path = self.dir / name
        path.write_bytes(build(sheets(text)))
        return wb.load_workbook(self.db, wb.import_workbook(self.db, path).workbook_id)

    def save(self, backlog_id: int, *cases, questions=("Which formats?",), locked=(), guidance="", number=7) -> int:
        request = generations.request_for_backlog(self.db, backlog_id, DesignerSettings(), user_guidance=guidance,
                                                  locked_test_cases=list(locked))
        outcome = Planner(request).on_result("first", CallResult(structured=response(*cases, number=number,
                                                                                    questions=list(questions)),
                                                                 subtype="success", model="claude-opus-5-5"))
        self.assertTrue(outcome.ok, outcome.validation_errors)
        return generations.save_outcome(self.db, backlog_id, outcome)


class HistoryTests(Base):
    def test_restore_and_versions(self):
        first = self.save(self.orders.id, case(1))
        second = self.save(self.orders.id, case(1), case(2, "corner_case"), guidance="More corner cases")
        self.assertEqual(content.current_run_id(self.db, self.orders.id), second)
        generations.restore_run(self.db, first)
        self.assertEqual(content.current_run_id(self.db, self.orders.id), first)
        self.assertEqual(len(content.current_cases(self.db, self.orders.id)), 1)
        self.assertEqual(generations.list_runs(self.db, self.orders.id)[0].user_guidance, "More corner cases")

        content.add_context(self.db, self.orders.id, content.DOMAIN_NOTE, "Managers only.")
        v2 = self.upload("Orders v2.xlsx", "Orders can be exported as CSV.")
        orders_v2 = v2.sheets[0].backlogs[1]
        self.assertEqual(v2.changes["changed"], ["Main|7|1"])
        # The new version shows the earlier version's test cases, runs and notes (linked by sheet + number).
        self.assertEqual(content.current_run_id(self.db, orders_v2.id), first)
        self.assertEqual([r.id for r in generations.list_runs(self.db, orders_v2.id)], [second, first])
        self.assertEqual(content.texts(self.db, orders_v2.id, content.DOMAIN_NOTE), ["Managers only."])
        third = self.save(orders_v2.id, case(1), case(2, "corner_case"), case(3, "corner_case"))
        self.assertEqual(content.current_run_id(self.db, orders_v2.id), third)
        self.assertEqual(content.current_run_id(self.db, self.orders.id), first)  # older versions stay as they were
        # Another backlog never shares them.
        self.assertIsNone(content.current_run_id(self.db, v2.sheets[0].backlogs[0].id))

    def test_only_valid_runs_can_be_restored(self):
        request = generations.request_for_backlog(self.db, self.orders.id, DesignerSettings())
        failed = Planner(request).on_result("first", CallResult(process_error="Claude Code could not be started."))
        run_id = generations.save_outcome(self.db, self.orders.id, failed)
        with self.assertRaises(ValueError):
            generations.restore_run(self.db, run_id)

    def test_regeneration_keeps_locked_cases(self):
        locked = case(3, "corner_case", title="Verify that a locked case stays")
        request = generations.request_for_backlog(self.db, self.orders.id, DesignerSettings(), locked_test_cases=[locked])
        planner = Planner(request)
        self.assertIn("Verify that a locked case stays", planner.first_call().content[-1]["text"])
        outcome = planner.on_result("first", CallResult(structured=response(case(1), case(2, "corner_case")),
                                                        subtype="success"))
        self.assertTrue(outcome.ok, outcome.validation_errors)
        cases = outcome.response["test_cases"]
        self.assertEqual([(c["id"], c["type"]) for c in cases],
                         [("TC-007-01", "happy_path"), ("TC-007-02", "corner_case"), ("TC-007-03", "corner_case")])
        self.assertEqual(cases[1]["title"], "Verify that a locked case stays")
        self.assertEqual(outcome.locked_ids, ["TC-007-02"])
        generations.save_outcome(self.db, self.orders.id, outcome)
        self.assertEqual([c.locked for c in content.current_cases(self.db, self.orders.id)], [False, True, False])


class ExportTests(Base):
    def setUp(self):
        super().setUp()
        self.save(self.orders.id, case(1), case(2, "corner_case", preconditions=[], steps=[
            "Open \"Orders\".", "Click \"Export\".", "Assert that a file is saved."]), questions=["Which formats?"])
        self.save(self.v1.sheets[0].backlogs[0].id, case(1, number=1, title="Verify that login works"),
                  questions=[], number=1)
        self.entries = content.export_entries(self.db, [b.id for b in self.v1.sheets[0].backlogs])

    def test_entries(self):
        self.assertEqual([(e.sheet, e.backlog_key, e.backlog_title, len(e.test_cases)) for e in self.entries],
                         [("Main", "1", "Login", 1), ("Main", "7", "Orders", 2)])
        self.assertEqual(self.entries[1].open_questions, ["Which formats?"])
        smoke = exporter.filter_by_tag(self.entries, "smoke")
        self.assertEqual([len(e.test_cases) for e in smoke], [1, 1])
        self.assertIs(exporter.filter_by_tag(self.entries, ""), self.entries)

    def test_json(self):
        one = exporter.export(self.dir / "one.json", "json", self.entries[1:], DesignerSettings())
        data = json.loads(one[0].read_text("utf-8"))
        self.assertEqual((data["backlog_number"], data["status"], len(data["test_cases"])), (7, "ok", 2))
        many = exporter.export(self.dir / "all.json", "json", self.entries, DesignerSettings())
        self.assertEqual([d["backlog_number"] for d in json.loads(many[0].read_text("utf-8"))], [1, 7])
        files = exporter.export(self.dir / "folder", "json", self.entries, DesignerSettings(), per_backlog_json=True)
        self.assertEqual([p.name for p in files], ["Main-backlog-1.json", "Main-backlog-7.json"])
        with self.assertRaises(ValueError):
            exporter.export(self.dir / "none.json", "json", [], DesignerSettings())

    def test_csv(self):
        [path] = exporter.export(self.dir / "cases.csv", "csv", self.entries, DesignerSettings())
        self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))  # UTF-8 with BOM
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        self.assertEqual(rows[0], ["id", "backlog_number", "backlog_title", "title", "type", "priority", "tags",
                                   "preconditions", "steps", "expected_result", "basis", "source_reference"])
        self.assertEqual(len(rows), 4)
        orders_corner = rows[3]
        self.assertEqual(orders_corner[2], "Orders")
        self.assertEqual(orders_corner[6], "regression")
        self.assertEqual(orders_corner[8], "1. Open \"Orders\".\n2. Click \"Export\".\n3. Assert that a file is saved.")
        self.assertEqual(rows[2][6], "smoke;regression")
        self.assertEqual(rows[2][7], "User is logged in")

    def test_columns_follow_the_preset(self):
        settings = DesignerSettings()
        settings.output_presets.append(OutputPreset(name="P", steps_style=PER_STEP_EXPECTED,
                                                    optional_fields={"basis": False, "source_reference": False}))
        settings.output_preset = "P"
        self.assertNotIn("basis", exporter.columns(settings))
        entry = exporter.ExportEntry("Main", "7", "Orders", [case(1, steps=[
            {"action": "Open it.", "expected": ""}, {"action": "Save.", "expected": "Assert that it is saved."}])])
        self.assertEqual(exporter.cell_value(entry.test_cases[0], "steps", entry),
                         "1. Open it.\n2. Save. => Assert that it is saved.")

    def test_xlsx(self):
        [path] = exporter.export(self.dir / "cases.xlsx", "xlsx", self.entries, DesignerSettings())
        book = load_workbook(path)
        self.assertEqual(book.sheetnames, ["Test Cases", "Open Questions"])
        cases, questions = book["Test Cases"], book["Open Questions"]
        self.assertEqual(cases.freeze_panes, "A2")
        self.assertEqual(cases["A1"].value, "id")
        self.assertEqual(cases.max_row, 4)
        self.assertTrue(cases["I2"].alignment.wrap_text)
        self.assertEqual([[c.value for c in row] for row in questions.iter_rows()],
                         [["sheet", "backlog_number", "question"], ["Main", "7", "Which formats?"]])


class WorkspaceTests(Base):
    def setUp(self):
        super().setUp()
        self.workspace = DesignerWorkspace(self.db, self.dir / "runs", None)
        self.workspace.export_dir = self.dir / "exports"
        self.workspace.tester_name = lambda: "Mohammad"
        self.workspace.resize(1400, 900)
        self.workspace.show()
        self.workspace.set_workbook(self.v1.id, "Orders v1.xlsx (v1)")
        FakeRun.created = []

    def tearDown(self):
        self.workspace.deleteLater()
        super().tearDown()

    def test_regenerate_history_filter_and_export(self):
        w = self.workspace
        self.save(self.orders.id, case(1), case(2, "corner_case"))
        w.select_backlog(self.orders.id)
        panel, results = w.backlog_panel, w.results_panel
        self.assertEqual(panel.generate_button.text(), "Regenerate test cases")
        self.assertEqual(panel.history_button.text(), "History (1)…")
        self.assertFalse(panel.keep_locked.isVisibleTo(panel))
        results.model.setData(results.model.index(1, len(results.model.columns) - 2), 2, 10)  # lock TC-007-02
        for _ in range(50):
            app.processEvents()
        self.assertTrue(panel.keep_locked.isVisibleTo(panel))
        self.assertIn("(1)", panel.keep_locked.text())

        # Tag filter: only the smoke-tagged happy path shows.
        index = results.tag_filter.findData("smoke")
        results.tag_filter.setCurrentIndex(index)
        self.assertEqual([c.data["id"] for c in results.visible_cases()], ["TC-007-01"])
        results.tag_filter.setCurrentIndex(0)

        panel.guidance.setText("Focus on permissions")
        FakeRun.answer = CallResult(structured=response(case(1, title="Verify that a new flow works")),
                                    subtype="success")
        with mock.patch.object(workspace_module, "GenerationRun", FakeRun):
            panel.generate_button.click()
            for _ in range(200):
                if w.run is None:
                    break
                app.processEvents()
        [run] = FakeRun.created
        self.assertEqual(run.request.user_guidance, "Focus on permissions")
        self.assertEqual([c["id"] for c in run.request.locked_test_cases], ["TC-007-02"])
        self.assertEqual(panel.guidance.text(), "")
        self.assertEqual([(c.data["title"], c.locked) for c in results.model.cases],
                         [("Verify that a new flow works", False), ("Verify that flow number 2 works", True)])
        self.assertEqual(panel.history_button.text(), "History (2)…")

        dialog = HistoryDialog(self.db, self.orders.id, "Orders", w)
        self.assertEqual(dialog.table.item(0, 7).text(), "current")
        self.assertEqual(dialog.table.item(0, 6).text(), "Focus on permissions")
        dialog.table.selectRow(1)
        self.assertTrue(dialog.restore_button.isEnabled())
        dialog.restore()
        self.assertEqual(dialog.table.item(1, 7).text(), "current")
        dialog.deleteLater()

        entries = content.export_entries(self.db, [self.orders.id])
        with mock.patch.object(QMessageBox, "exec", return_value=0):
            written = w.write_export(self.dir / "out.xlsx", "xlsx", entries, DesignerSettings(), False, "")
        self.assertTrue(written[0].is_file())
        with self.db.Session() as orm:
            log = orm.scalars(select(ExportLog)).one()
            self.assertEqual((log.user, log.format, log.selection["backlogs"]), ("Mohammad", "xlsx", ["Main|7"]))

    def test_export_dialog(self):
        dialog = ExportDialog("#7 Orders", 2, ["smoke", "regression"])
        self.assertTrue(dialog.this_radio.isChecked())
        self.assertFalse(dialog.per_backlog.isEnabled())
        dialog.format_combo.setCurrentIndex(dialog.format_combo.findData("json"))
        dialog.all_radio.setChecked(True)
        self.assertTrue(dialog.per_backlog.isEnabled())
        dialog.tag_combo.setCurrentIndex(1)
        self.assertEqual((dialog.scope(), dialog.fmt(), dialog.tag()), ("workbook", "json", "smoke"))
        none_shown = ExportDialog(None, 1, [])
        self.assertTrue(none_shown.all_radio.isChecked())
        self.assertFalse(none_shown.this_radio.isEnabled())
        dialog.deleteLater()
        none_shown.deleteLater()


if __name__ == "__main__":
    unittest.main()
