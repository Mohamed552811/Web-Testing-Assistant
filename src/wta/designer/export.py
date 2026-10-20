"""Export test cases (spec 8.1 (9)): JSON, CSV and Excel, for one backlog or several.

- JSON: the response object per backlog (spec 5.2), in one file per backlog or one array.
- CSV: UTF-8 with BOM, one row per case; tags joined with ";", preconditions and steps with newlines,
  steps numbered "1. ", "2. " … on export only.
- Excel: sheet "Test Cases" (same columns, frozen header, wrapped text) and sheet "Open Questions".
Columns follow the selected output preset.
"""

import csv
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wta.designer import prompt_builder as pb
from wta.designer.settings import DesignerSettings

FORMATS = {"json": "JSON (.json)", "csv": "CSV (.csv)", "xlsx": "Excel (.xlsx)"}
ALL_COLUMNS = ("id", "backlog_number", "backlog_title", "title", "type", "priority", "tags", "preconditions", "steps",
               "expected_result", "basis", "source_reference")
STEP_SEPARATOR = " => "
QUESTION_COLUMNS = ("sheet", "backlog_number", "question")
_COLUMN_WIDTHS = {"id": 12, "backlog_number": 10, "backlog_title": 28, "title": 45, "type": 13, "priority": 10,
                  "tags": 18, "preconditions": 40, "steps": 70, "expected_result": 45, "basis": 11,
                  "source_reference": 24}


@dataclass
class ExportEntry:
    """The current test cases and open questions of one backlog."""

    sheet: str
    backlog_key: str
    backlog_title: str
    test_cases: list[dict]
    open_questions: list[str] = field(default_factory=list)

    def response(self) -> dict:
        """The spec 5.2 response object."""
        return {"backlog_number": pb.backlog_number_value(self.backlog_key), "backlog_title": self.backlog_title,
                "status": "ok", "test_cases": self.test_cases, "open_questions": self.open_questions}


def columns(settings: DesignerSettings) -> list[str]:
    preset = settings.preset()
    return [c for c in ALL_COLUMNS if preset.includes(c)]


def filter_by_tag(entries: list[ExportEntry], tag: str | None) -> list[ExportEntry]:
    if not tag:
        return entries
    return [ExportEntry(e.sheet, e.backlog_key, e.backlog_title, [c for c in e.test_cases if tag in (c.get("tags") or [])],
                        e.open_questions) for e in entries]


def _step_text(step) -> str:
    if isinstance(step, dict):
        expected = step.get("expected", "")
        return step.get("action", "") + (f"{STEP_SEPARATOR}{expected}" if expected else "")
    return str(step)


def cell_value(case: dict, column: str, entry: ExportEntry) -> str:
    if column == "backlog_title":
        return entry.backlog_title
    if column == "backlog_number":
        return str(case.get("backlog_number", entry.backlog_key))
    value = case.get(column)
    if column == "tags":
        return ";".join(value or [])
    if column == "preconditions":
        return "\n".join(value or [])
    if column == "steps":
        return "\n".join(f"{n}. {_step_text(s)}" for n, s in enumerate(value or [], 1))
    return "" if value is None else str(value)


def rows(entries: list[ExportEntry], settings: DesignerSettings) -> list[list[str]]:
    names = columns(settings)
    return [[cell_value(case, name, entry) for name in names] for entry in entries for case in entry.test_cases]


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_") or "backlog"


def write_json(path: Path, entries: list[ExportEntry], per_backlog: bool = False) -> list[Path]:
    """One file holding the object (one backlog) or an array (several); or a folder with one file per backlog."""
    if per_backlog:
        path.mkdir(parents=True, exist_ok=True)
        written = []
        for entry in entries:
            target = path / f"{safe_name(entry.sheet)}-backlog-{safe_name(entry.backlog_key)}.json"
            target.write_text(json.dumps(entry.response(), indent=2, ensure_ascii=False), encoding="utf-8")
            written.append(target)
        return written
    data = entries[0].response() if len(entries) == 1 else [e.response() for e in entries]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return [path]


def write_csv(path: Path, entries: list[ExportEntry], settings: DesignerSettings) -> list[Path]:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns(settings))
        writer.writerows(rows(entries, settings))
    return [path]


def write_xlsx(path: Path, entries: list[ExportEntry], settings: DesignerSettings) -> list[Path]:
    book = Workbook()
    sheet = book.active
    sheet.title = "Test Cases"
    names = columns(settings)
    sheet.append(names)
    for row in rows(entries, settings):
        sheet.append(row)
    questions = book.create_sheet("Open Questions")
    questions.append(list(QUESTION_COLUMNS))
    for entry in entries:
        for question in entry.open_questions:
            questions.append([entry.sheet, entry.backlog_key, question])
    header_fill = PatternFill("solid", fgColor="ECEEF2")
    for ws, widths in ((sheet, [_COLUMN_WIDTHS[n] for n in names]), (questions, [18, 14, 100])):
        ws.freeze_panes = "A2"
        for index, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return [path]


def export(path: Path, fmt: str, entries: list[ExportEntry], settings: DesignerSettings,
           per_backlog_json: bool = False) -> list[Path]:
    if not entries:
        raise ValueError("There are no test cases to export.")
    if fmt == "json":
        return write_json(path, entries, per_backlog_json)
    if fmt == "csv":
        return write_csv(path, entries, settings)
    if fmt == "xlsx":
        return write_xlsx(path, entries, settings)
    raise ValueError(f"Unknown export format {fmt!r}")
