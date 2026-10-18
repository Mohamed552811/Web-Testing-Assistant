"""Word report (python-docx)."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from wta.export.html import STATUS_COLORS, other_screenshot_rows
from wta.export.model import FINAL_COLUMNS, Report, summary_pairs

RESULT_HEADERS = ("Type", "Scenario / page", "#", "Action / input", "Expected", "Actual", "Claude", "Final",
                  "Review comment")


def _shade(cell, text: str) -> None:
    color = STATUS_COLORS.get(text.split(" ")[0])
    if color:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), color.lstrip("#"))
        cell._tc.get_or_add_tcPr().append(shading)


def _key_value(doc, headers, rows):
    """A two-column Field / Value table with a narrow first column."""
    table = _table(doc, headers, rows)
    table.autofit = False
    for row in table.rows:
        row.cells[0].width, row.cells[1].width = Cm(4.5), Cm(20)
    return table


def _table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, headers):
        cell.text = name
        cell.paragraphs[0].runs[0].bold = True
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value if value is not None else "")
    return table


def write_docx(report: Report, path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(9)
    picture_width = min(Cm(16), section.page_width - section.left_margin - section.right_margin)

    doc.add_heading(report.title, level=0)
    _key_value(doc, ("Item", "Value"), summary_pairs(report))

    doc.add_heading("Final status by type", level=1)
    _table(doc, ("Type",) + FINAL_COLUMNS + ("Total",), [
        (kind, *(counter.get(c, 0) for c in FINAL_COLUMNS), sum(counter.values()))
        for kind, counter in report.counts.items()
    ])

    if report.scenarios:
        doc.add_heading("Scenarios", level=1)
        _table(doc, ("Batch", "Scenario", "Steps", "Expected result", "Test data"), [
            (info.batch, info.name, info.numbered_steps(), info.expected_result, info.test_data)
            for info in report.scenarios
        ])

    doc.add_heading("Results", level=1)
    table = _table(doc, RESULT_HEADERS, [
        (f"{s.kind} {s.batch}", s.scenario, s.order, s.action, s.expected, s.actual, s.claude_status,
         s.final_status, s.comment)
        for s in report.rows
    ])
    for row, step in zip(table.rows[1:], report.rows):
        _shade(row.cells[6], step.claude_status)
        _shade(row.cells[7], step.final_status)

    doc.add_heading(f"Bugs ({len(report.bugs)})", level=1)
    if not report.bugs:
        doc.add_paragraph("No bugs in this report.")
    for i, bug in enumerate(report.bugs, start=1):
        doc.add_heading(f"{i}. {bug.title}", level=2)
        _key_value(doc, ("Field", "Value"), [
            ("Severity", bug.severity), ("Priority", bug.priority), ("Review state", bug.review_state),
            ("Where", f"{bug.kind}: {bug.scenario}, step {bug.step_order}"),
            ("Repro steps", bug.tfs_repro_steps), ("Bug report", bug.report_text),
        ])
        for shot in bug.screenshots:
            doc.add_picture(str(shot), width=picture_width)

    others = other_screenshot_rows(report)
    if others:
        doc.add_heading("Other screenshots", level=1)
        for step in others:
            doc.add_heading(f"{step.scenario}, step {step.order}: {step.final_status}", level=3)
            for shot in step.screenshots:
                doc.add_picture(str(shot), width=picture_width)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
