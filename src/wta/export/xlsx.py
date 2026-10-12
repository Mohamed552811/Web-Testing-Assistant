"""Excel report: Summary, Results and Bugs sheets, with screenshots embedded."""

import struct
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from wta.export.model import FINAL_COLUMNS, Report, summary_pairs

THUMB_WIDTH = 320  # px
HEADER_FILL = PatternFill("solid", fgColor="203040")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILLS = {
    "Pass": "C6EFCE", "Fail": "FFC7CE", "Bug": "FFC7CE", "Blocked": "FFEB9C", "Not run": "E7E6E6",
}
WRAP = Alignment(wrap_text=True, vertical="top")

RESULT_COLUMNS = (
    ("Type", 11), ("Batch", 7), ("Scenario / page", 26), ("#", 5), ("Screen", 16), ("Field / element", 20),
    ("Action / input", 34), ("Expected", 30), ("Actual", 34), ("Claude status", 12), ("Final status", 16),
    ("Review", 11), ("Review comment", 26), ("Reviewed at", 16), ("Claude reasoning", 40), ("Page URL", 30),
    ("Console errors", 30), ("Screenshot", 46),
)
# Column names follow Azure DevOps / TFS Bug work items where one exists, so the sheet can be imported.
BUG_COLUMNS = (
    ("Work Item Type", 14), ("Title", 40), ("Repro Steps", 60), ("Severity", 12), ("Priority", 9), ("Tags", 20),
    ("Review state", 16), ("Type", 11), ("Scenario / page", 26), ("Step", 6), ("Bug report", 60), ("Screenshot", 46),
)


def png_size(path: Path) -> tuple[int, int] | None:
    """Width and height from a PNG header, or None if the file is not a PNG."""
    with open(path, "rb") as f:
        head = f.read(24)
    if len(head) < 24 or head[:8] != b"\x89PNG\r\n\x1a\n" or head[12:16] != b"IHDR":
        return None
    return struct.unpack(">II", head[16:24])


class PngImage(Image):
    """openpyxl Image for a PNG file without Pillow (openpyxl only needs the size and the bytes).

    Relies on openpyxl's Image._data(); openpyxl is pinned to 3.1.x in pyproject.toml.
    """

    def __init__(self, path: Path, width: int, height: int) -> None:  # noqa: super().__init__ needs Pillow
        self.ref = str(path)
        self.format = "png"
        self.width, self.height = width, height

    def _data(self) -> bytes:
        return Path(self.ref).read_bytes()


def _embed(ws, path: Path, cell: str) -> float | None:
    """Embed a screenshot scaled to THUMB_WIDTH at `cell`. Returns the row height needed (points)."""
    size = png_size(path)
    if size is None:
        return None
    width, height = size
    scale = min(1.0, THUMB_WIDTH / width) if width else 1.0
    image = PngImage(path, int(width * scale), int(height * scale))
    ws.add_image(image, cell)
    return image.height * 0.75 + 6  # px to points, with a little margin


def _print_setup(ws) -> None:
    """Wide sheets print landscape and fit to one page wide."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def _header(ws, columns) -> None:
    _print_setup(ws)
    for index, (name, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=index, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def _fill_status(cell, text: str) -> None:
    color = STATUS_FILLS.get(text.split(" ")[0])
    if color:
        cell.fill = PatternFill("solid", fgColor=color)


def write_xlsx(report: Report, path: Path) -> Path:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = report.title
    summary["A1"].font = Font(bold=True, size=14)
    for row, (key, value) in enumerate(summary_pairs(report), start=3):
        summary.cell(row=row, column=1, value=key).font = Font(bold=True)
        summary.cell(row=row, column=2, value=value)
    row = len(summary_pairs(report)) + 5
    summary.cell(row=row, column=1, value="Final status by type").font = Font(bold=True, size=12)
    row += 1
    for col, name in enumerate(("Type",) + FINAL_COLUMNS + ("Total",), start=1):
        cell = summary.cell(row=row, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for kind, counter in report.counts.items():
        row += 1
        summary.cell(row=row, column=1, value=kind)
        for col, key in enumerate(FINAL_COLUMNS, start=2):
            summary.cell(row=row, column=col, value=counter.get(key, 0))
        summary.cell(row=row, column=len(FINAL_COLUMNS) + 2, value=sum(counter.values()))
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 60
    for col in range(3, len(FINAL_COLUMNS) + 3):
        summary.column_dimensions[get_column_letter(col)].width = 15

    results = wb.create_sheet("Results")
    _header(results, RESULT_COLUMNS)
    shot_col = get_column_letter(len(RESULT_COLUMNS))
    for r, step in enumerate(report.rows, start=2):
        values = (
            step.kind, step.batch, step.scenario, step.order, step.screen, step.element, step.action,
            step.expected, step.actual, step.claude_status, step.final_status, step.decision, step.comment,
            step.reviewed_at, step.reasoning, step.page_url, "\n".join(step.console_errors),
            "\n".join(str(p) for p in step.screenshots[1:]),
        )
        for c, value in enumerate(values, start=1):
            results.cell(row=r, column=c, value=value).alignment = WRAP
        _fill_status(results.cell(row=r, column=10), step.claude_status)
        _fill_status(results.cell(row=r, column=11), step.final_status)
        if step.screenshots:
            height = _embed(results, step.screenshots[0], f"{shot_col}{r}")
            if height:
                results.row_dimensions[r].height = max(height, 30)

    bugs = wb.create_sheet("Bugs")
    _header(bugs, BUG_COLUMNS)
    shot_col = get_column_letter(len(BUG_COLUMNS))
    for r, bug in enumerate(report.bugs, start=2):
        values = (
            "Bug", bug.title, bug.tfs_repro_steps, bug.severity, bug.priority, "WebTestingAssistant",
            bug.review_state, bug.kind, bug.scenario, bug.step_order, bug.report_text, "",
        )
        for c, value in enumerate(values, start=1):
            bugs.cell(row=r, column=c, value=value).alignment = WRAP
        if bug.screenshots:
            height = _embed(bugs, bug.screenshots[0], f"{shot_col}{r}")
            if height:
                bugs.row_dimensions[r].height = max(height, 30)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
