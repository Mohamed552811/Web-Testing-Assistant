"""Read test cases from a CSV or Excel file into scenarios for the library.

Two layouts are understood, by their column names (case and spacing do not matter):
- one row per test case: Name/Title, Steps (one per line), Expected result, Test data;
- a title row followed by one row per step, as Azure DevOps exports test cases:
  Title, then rows with Step Action and Step Expected (and an empty Title).
"""

import csv
import html
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from wta.db.scenarios import ScenarioInput

MAX_CASES = 500
NAME_COLUMNS = ("name", "title", "scenario", "test case", "test case title", "summary")
STEP_COLUMNS = ("steps", "test steps", "step action", "action", "step", "actions")
EXPECTED_COLUMNS = ("expected result", "expected results", "expected", "step expected", "expected outcome")
DATA_COLUMNS = ("test data", "data")


@dataclass
class ImportResult:
    cases: list[ScenarioInput] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _clean(value) -> str:
    """Cell text without HTML (Azure DevOps exports steps as HTML) and with tidy whitespace."""
    text = "" if value is None else str(value)
    if "<" in text and ">" in text:
        text = re.sub(r"(?i)<br\s*/?>|</p>|</div>", "\n", text)
        text = re.sub(r"<[^>]+>", "", text)
    text = html.unescape(text).replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _header_key(value) -> str:
    return re.sub(r"\s+", " ", _clean(value).lower())


def read_rows(path: Path) -> list[list[str]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook  # already a dependency (Excel export)
        book = load_workbook(path, read_only=True, data_only=True)
        try:  # a read-only workbook keeps the file open (and locked on Windows) until closed
            return [[_clean(cell) for cell in row] for row in book.worksheets[0].iter_rows(values_only=True)]
        finally:
            book.close()
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    # newline="" keeps line breaks inside quoted cells (multi-line steps)
    return [[_clean(cell) for cell in row] for row in csv.reader(io.StringIO(text, newline=""), dialect)]


def _find(headers: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in headers:
            return headers.index(name)
    return None


def parse_rows(rows: list[list[str]]) -> ImportResult:
    result = ImportResult()
    rows = [[_clean(cell) for cell in row] for row in rows]
    rows = [row for row in rows if any(cell for cell in row)]
    if not rows:
        result.warnings.append("The file is empty.")
        return result
    headers = [_header_key(h) for h in rows[0]]
    name_col, steps_col = _find(headers, NAME_COLUMNS), _find(headers, STEP_COLUMNS)
    expected_col, data_col = _find(headers, EXPECTED_COLUMNS), _find(headers, DATA_COLUMNS)
    if name_col is None:
        result.warnings.append("No Name or Title column found in the first row.")
        return result

    def cell(row: list[str], col: int | None) -> str:
        return row[col] if col is not None and col < len(row) else ""

    cases: list[dict] = []
    for number, row in enumerate(rows[1:], start=2):
        name, step, expected = cell(row, name_col), cell(row, steps_col), cell(row, expected_col)
        if name:
            cases.append({"name": name, "steps": [], "expected": [], "data": cell(row, data_col)})
            if step:
                cases[-1]["steps"].extend(line for line in step.split("\n") if line)
            if expected:
                cases[-1]["expected"].append(expected)
        elif step and cases:  # a step row of the case above (Azure DevOps layout)
            cases[-1]["steps"].append(" ".join(step.split("\n")) + (f" (expect: {expected})" if expected else ""))
            if expected:
                cases[-1]["step_expected"] = expected
        elif step or expected:
            result.warnings.append(f"Row {number} has a step but no test case above it; skipped.")
    for case in cases[:MAX_CASES]:
        expected = "\n".join(case["expected"]) or case.get("step_expected", "")
        steps = case["steps"]
        if not case["expected"] and steps and case.get("step_expected"):
            # The last step's expectation becomes the expected result instead of staying inline.
            steps = steps[:-1] + [re.sub(r" \(expect: .*\)$", "", steps[-1])]
        try:
            result.cases.append(ScenarioInput(case["name"][:200], "\n".join(steps), expected, case["data"]).cleaned())
        except ValueError as exc:
            result.warnings.append(f"“{case['name']}”: {exc}")
    if len(cases) > MAX_CASES:
        result.warnings.append(f"Only the first {MAX_CASES} test cases were read.")
    if not result.cases and not result.warnings:
        result.warnings.append("No test cases found.")
    return result


def import_file(path: str | Path) -> ImportResult:
    path = Path(path)
    try:
        return parse_rows(read_rows(path))
    except Exception as exc:  # a broken or unexpected file: tell the user, do not crash
        return ImportResult(warnings=[f"Could not read {path.name}: {exc}"])
