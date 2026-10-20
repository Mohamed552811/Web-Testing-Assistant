"""Manual Smoke and Regression checklist runs (spec 8.6): made from designer test cases, followed by a tester.

Each item keeps a snapshot of its test case, so later edits in the designer do not change a run; while a run is
still a draft (nothing recorded yet), "refresh from designer" takes the current versions again.
"""

import csv
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from wta.db import Database
from wta.db.designer_models import ChecklistItem, ChecklistRun, DesignerTestCase
from wta.db.models import utcnow

SMOKE, REGRESSION = "smoke", "regression"
MODES = (SMOKE, REGRESSION)
RESULTS = {"pass": "Pass", "fail": "Fail", "blocked": "Blocked", "skipped": "Skipped"}
DRAFT, IN_PROGRESS, COMPLETED = "draft", "in_progress", "completed"
STATUS_LABELS = {DRAFT: "Draft", IN_PROGRESS: "In progress", COMPLETED: "Completed"}
EXPORT_COLUMNS = ("position", "id", "title", "preconditions", "steps", "expected_result", "result", "notes",
                  "attachment", "executed_at")


@dataclass
class HandOffCase:
    """A test case to hand over: the designer row (for refresh) and its data."""

    row_id: int
    backlog_id: int
    data: dict


@dataclass
class ItemView:
    id: int
    position: int
    test_case_ref: str
    snapshot: dict
    result: str | None
    notes: str
    attachment: str | None
    executed_at: datetime | None


@dataclass
class RunView:
    id: int
    mode: str
    name: str
    created_at: datetime
    status: str
    source: dict
    items: list[ItemView] = field(default_factory=list)

    def counts(self) -> dict[str, int]:
        counts = {key: 0 for key in RESULTS}
        counts["not_run"] = 0
        for item in self.items:
            counts[item.result or "not_run"] += 1
        return counts

    @property
    def progress_text(self) -> str:
        c = self.counts()
        done = len(self.items) - c["not_run"]
        return (f"{done} of {len(self.items)} done · Pass {c['pass']} · Fail {c['fail']} · "
                f"Blocked {c['blocked']} · Skipped {c['skipped']}")


def default_name(mode: str, workbook: str, now: datetime | None = None) -> str:
    """Spec 8.6: "Smoke – <workbook> – <date>"."""
    return f"{mode.capitalize()} – {Path(workbook).stem} – {(now or datetime.now()):%Y-%m-%d}"


def _item_view(item: ChecklistItem) -> ItemView:
    return ItemView(item.id, item.position, item.test_case_ref, dict(item.snapshot), item.result, item.notes or "",
                    item.attachment_url, item.executed_at)


def _run_view(run: ChecklistRun) -> RunView:
    return RunView(run.id, run.mode, run.name, run.created_at, run.status, dict(run.source or {}),
                   [_item_view(i) for i in run.items])


def create_run(db: Database, mode: str, name: str, cases: list[HandOffCase], source: dict) -> int:
    if mode not in MODES:
        raise ValueError(f"Unknown mode {mode!r}")
    if not cases:
        raise ValueError("Choose at least one test case.")
    name = name.strip()[:200] or "Checklist run"
    with db.Session() as orm, orm.begin():
        run = ChecklistRun(mode=mode, name=name, source=source)
        orm.add(run)
        orm.flush()
        _append(run, cases)
        return run.id


def add_to_run(db: Database, run_id: int, cases: list[HandOffCase]) -> int:
    """Add test cases to an open run; a case already in it (same designer row) is not added twice."""
    with db.Session() as orm, orm.begin():
        run = orm.get(ChecklistRun, run_id)
        if run is None or run.status == COMPLETED:
            raise ValueError("That run is not open any more.")
        known = {i.test_case_ref for i in run.items}
        return _append(run, [c for c in cases if str(c.row_id) not in known])


def _append(run: ChecklistRun, cases: list[HandOffCase]) -> int:
    start = len(run.items)
    for offset, case in enumerate(cases, 1):
        run.items.append(ChecklistItem(test_case_ref=str(case.row_id), backlog_item_id=case.backlog_id,
                                       position=start + offset, snapshot=dict(case.data)))
    return len(cases)


def list_runs(db: Database, mode: str | None = None, include_completed: bool = False) -> list[RunView]:
    with db.Session() as orm:
        query = select(ChecklistRun).order_by(ChecklistRun.id.desc())
        if mode:
            query = query.where(ChecklistRun.mode == mode)
        if not include_completed:
            query = query.where(ChecklistRun.status != COMPLETED)
        return [_run_view(r) for r in orm.scalars(query)]


def open_count(db: Database, mode: str) -> int:
    return len(list_runs(db, mode))


def load_run(db: Database, run_id: int) -> RunView:
    with db.Session() as orm:
        run = orm.get(ChecklistRun, run_id)
        if run is None:
            raise LookupError(f"Checklist run {run_id} not found")
        return _run_view(run)


def record(db: Database, item_id: int, result: str | None, notes: str | None = None) -> None:
    """Pass / Fail / Blocked / Skipped (None clears it) and notes. The run moves from draft to in progress."""
    if result is not None and result not in RESULTS:
        raise ValueError(f"Unknown result {result!r}")
    with db.Session() as orm, orm.begin():
        item = orm.get(ChecklistItem, item_id)
        if item is None:
            raise LookupError(f"Checklist item {item_id} not found")
        if item.run.status == COMPLETED:
            raise ValueError("The run is completed; reopen it to change results.")
        item.result = result
        item.executed_at = utcnow() if result else None
        if notes is not None:
            item.notes = notes.strip()
        if item.run.status == DRAFT and (result or item.notes):
            item.run.status = IN_PROGRESS


def attach(db: Database, item_id: int, file: Path, attachments_dir: Path) -> str:
    """Copy a screenshot into the app's attachments folder and link it to the item."""
    file = Path(file)
    target_dir = Path(attachments_dir) / f"run-{_run_id_of(db, item_id)}"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"item-{item_id}{file.suffix.lower()}"
    shutil.copyfile(file, target)
    relative = target.relative_to(attachments_dir).as_posix()
    with db.Session() as orm, orm.begin():
        item = orm.get(ChecklistItem, item_id)
        item.attachment_url = relative
        if item.run.status == DRAFT:
            item.run.status = IN_PROGRESS
    return relative


def _run_id_of(db: Database, item_id: int) -> int:
    with db.Session() as orm:
        item = orm.get(ChecklistItem, item_id)
        if item is None:
            raise LookupError(f"Checklist item {item_id} not found")
        return item.run_id


def set_completed(db: Database, run_id: int, completed: bool) -> None:
    with db.Session() as orm, orm.begin():
        run = orm.get(ChecklistRun, run_id)
        if run is None:
            raise LookupError(f"Checklist run {run_id} not found")
        if completed:
            run.status, run.completed_at = COMPLETED, utcnow()
        else:
            run.status = IN_PROGRESS if any(i.result or i.notes for i in run.items) else DRAFT
            run.completed_at = None


def refresh_from_designer(db: Database, run_id: int) -> int:
    """Draft runs only: replace each snapshot with the designer's current version. Returns how many changed."""
    with db.Session() as orm, orm.begin():
        run = orm.get(ChecklistRun, run_id)
        if run is None:
            raise LookupError(f"Checklist run {run_id} not found")
        if run.status != DRAFT:
            raise ValueError("Only a draft run (no results yet) can be refreshed from the designer.")
        changed = 0
        for item in run.items:
            row = orm.get(DesignerTestCase, int(item.test_case_ref)) if item.test_case_ref.isdigit() else None
            if row is not None and row.deleted_at is None and row.data != item.snapshot:
                item.snapshot = dict(row.data)
                changed += 1
        return changed


def delete_run(db: Database, run_id: int, attachments_dir: Path | None = None) -> None:
    with db.Session() as orm, orm.begin():
        run = orm.get(ChecklistRun, run_id)
        if run is not None:
            orm.delete(run)
    if attachments_dir is not None:
        shutil.rmtree(Path(attachments_dir) / f"run-{run_id}", ignore_errors=True)


# --- export -----------------------------------------------------------------------------------------


def _step(step) -> str:
    if isinstance(step, dict):
        return step.get("action", "") + (f" => {step['expected']}" if step.get("expected") else "")
    return str(step)


def export_rows(run: RunView) -> list[list[str]]:
    rows = []
    for item in run.items:
        s = item.snapshot
        rows.append([str(item.position), s.get("id", ""), s.get("title", ""), "\n".join(s.get("preconditions") or []),
                     "\n".join(f"{n}. {_step(step)}" for n, step in enumerate(s.get("steps") or [], 1)),
                     s.get("expected_result", ""), RESULTS.get(item.result, "Not run"), item.notes,
                     item.attachment or "", item.executed_at.strftime("%Y-%m-%d %H:%M") if item.executed_at else ""])
    return rows


def export_run(run: RunView, path: Path, fmt: str) -> Path:
    """Spec 8.6: the run can be exported to Excel or CSV."""
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = export_rows(run)
    if fmt == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(EXPORT_COLUMNS)
            writer.writerows(rows)
        return path
    if fmt != "xlsx":
        raise ValueError(f"Unknown export format {fmt!r}")
    book = Workbook()
    sheet = book.active
    sheet.title = "Checklist"
    sheet.append(list(EXPORT_COLUMNS))
    for row in rows:
        sheet.append(row)
    summary = book.create_sheet("Summary")
    summary.append(["Run", run.name])
    summary.append(["Mode", run.mode])
    summary.append(["Status", STATUS_LABELS.get(run.status, run.status)])
    for key, count in run.counts().items():
        summary.append([RESULTS.get(key, "Not run"), count])
    for index, width in enumerate((9, 12, 45, 40, 70, 45, 10, 40, 30, 17), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    book.save(path)
    return path
